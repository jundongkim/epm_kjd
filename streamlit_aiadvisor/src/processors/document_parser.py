"""
다양한 타입의 문서를 파싱하는 모듈
"""
import os
import json
from typing import List, Dict, Any, Optional
from datetime import datetime

import pdfplumber
from docx import Document
from pptx import Presentation
import pandas as pd

from langchain_community.document_loaders import (
    PyPDFLoader, 
    UnstructuredWordDocumentLoader, 
    UnstructuredPowerPointLoader,
    TextLoader
)
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import DirectoryLoader

from src.utils.config import DATA_DIR

class DocumentParser:
    """다양한 타입의 문서를 파싱하는 클래스"""

    def __init__(self, file_path: Optional[str] = None):
        """
        파서 초기화
        
        Args:
            file_path (str, optional): 파싱할 파일 경로. 없으면 DATA_DIR 전체를 대상으로 함
        """
        self.file_path = file_path
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len,
        )
    
    def _extract_metadata_from_filename(self, file_path: str) -> Dict[str, Any]:
        """
        파일명에서 메타데이터 추출
        
        Args:
            file_path (str): 파일 경로
            
        Returns:
            Dict[str, Any]: 추출된 메타데이터
        """
        file_name = os.path.basename(file_path)
        file_ext = os.path.splitext(file_name)[1].lower()
        file_stats = os.stat(file_path)
        mod_date = datetime.fromtimestamp(file_stats.st_mtime).strftime("%Y-%m-%d")
        
        return {
            "filename": file_name,
            "file_extension": file_ext,
            "file_path": file_path,
            "modified_date": mod_date,
            "file_size_kb": round(file_stats.st_size / 1024, 2)
        }
    
    def parse_pdf(self, file_path: str) -> Dict[str, Any]:
        """
        PDF 파일을 파싱
        
        Args:
            file_path (str): PDF 파일 경로
            
        Returns:
            Dict[str, Any]: 파싱 결과와 메타데이터
        """
        metadata = self._extract_metadata_from_filename(file_path)
        content = []
        
        try:
            loader = PyPDFLoader(file_path)
            pages = loader.load()
            
            # 페이지별 내용 추출
            for page in pages:
                content.append(page.page_content)
            
            metadata["content"] = content
            metadata["page_count"] = len(pages)
            
            # 텍스트 추출이 제대로 이루어지지 않았을 경우 pdfplumber로 다시 시도
            if not any(content) or all(len(page.strip()) == 0 for page in content):
                with pdfplumber.open(file_path) as pdf:
                    content = [page.extract_text() for page in pdf.pages if page.extract_text()]
                metadata["content"] = content
            
            return metadata
        
        except Exception as e:
            print(f"PDF 파싱 오류 ({file_path}): {str(e)}")
            metadata["content"] = ["[파싱 오류 발생]"]
            metadata["error"] = str(e)
            return metadata
    
    def parse_docx(self, file_path: str) -> Dict[str, Any]:
        """
        Word 문서 파싱
        
        Args:
            file_path (str): Word 파일 경로
            
        Returns:
            Dict[str, Any]: 파싱 결과와 메타데이터
        """
        metadata = self._extract_metadata_from_filename(file_path)
        content = []
        
        try:
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content.extend(paragraphs)
            
            # 테이블 내용 추출
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells if cell.text.strip()])
                    if row_text.strip():
                        content.append(row_text)
            
            metadata["content"] = content
            metadata["paragraph_count"] = len(paragraphs)
            
            return metadata
        
        except Exception as e:
            print(f"Word 파싱 오류 ({file_path}): {str(e)}")
            
            try:
                # 보조 로더 사용
                loader = UnstructuredWordDocumentLoader(file_path)
                doc_content = loader.load()
                metadata["content"] = [doc.page_content for doc in doc_content]
                return metadata
            except Exception as e2:
                metadata["content"] = ["[파싱 오류 발생]"]
                metadata["error"] = f"{str(e)} / {str(e2)}"
                return metadata
    
    def parse_pptx(self, file_path: str) -> Dict[str, Any]:
        """
        PowerPoint 파일 파싱
        
        Args:
            file_path (str): PowerPoint 파일 경로
            
        Returns:
            Dict[str, Any]: 파싱 결과와 메타데이터
        """
        metadata = self._extract_metadata_from_filename(file_path)
        content = []
        
        try:
            pres = Presentation(file_path)
            
            for i, slide in enumerate(pres.slides):
                slide_content = []
                
                # 제목 추출
                if slide.shapes.title and slide.shapes.title.has_text_frame:
                    title_text = slide.shapes.title.text.strip()
                    if title_text:
                        slide_content.append(f"Slide {i+1} Title: {title_text}")
                
                # 텍스트 상자 내용 추출
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content.append(shape.text.strip())
                
                if slide_content:
                    content.append("\n".join(slide_content))
            
            metadata["content"] = content
            metadata["slide_count"] = len(pres.slides)
            
            return metadata
        
        except Exception as e:
            print(f"PowerPoint 파싱 오류 ({file_path}): {str(e)}")
            
            try:
                # 보조 로더 사용
                loader = UnstructuredPowerPointLoader(file_path)
                ppt_content = loader.load()
                metadata["content"] = [doc.page_content for doc in ppt_content]
                return metadata
            except Exception as e2:
                metadata["content"] = ["[파싱 오류 발생]"]
                metadata["error"] = f"{str(e)} / {str(e2)}"
                return metadata
    
    def parse_txt(self, file_path: str) -> Dict[str, Any]:
        """
        텍스트 파일 파싱
        
        Args:
            file_path (str): 텍스트 파일 경로
            
        Returns:
            Dict[str, Any]: 파싱 결과와 메타데이터
        """
        metadata = self._extract_metadata_from_filename(file_path)
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.readlines()
            
            # 빈 줄 제거
            content = [line.strip() for line in content if line.strip()]
            metadata["content"] = content
            metadata["line_count"] = len(content)
            
            return metadata
        
        except UnicodeDecodeError:
            # UTF-8로 디코딩 실패 시 다른 인코딩 시도
            try:
                with open(file_path, 'r', encoding='cp949') as f:
                    content = f.readlines()
                content = [line.strip() for line in content if line.strip()]
                metadata["content"] = content
                metadata["line_count"] = len(content)
                return metadata
            except Exception as e:
                metadata["content"] = ["[파싱 오류 발생]"]
                metadata["error"] = str(e)
                return metadata
        except Exception as e:
            print(f"텍스트 파일 파싱 오류 ({file_path}): {str(e)}")
            metadata["content"] = ["[파싱 오류 발생]"]
            metadata["error"] = str(e)
            return metadata
    
    def parse_markdown(self, file_path: str) -> Dict[str, Any]:
        """
        마크다운 파일 파싱
        
        Args:
            file_path (str): 마크다운 파일 경로
            
        Returns:
            Dict[str, Any]: 파싱 결과와 메타데이터
        """
        metadata = self._extract_metadata_from_filename(file_path)
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.readlines()
            
            # 빈 줄 제거
            content = [line.strip() for line in content if line.strip()]
            
            # 마크다운 구조 분석 (헤더, 목록 등)
            headers = []
            for line in content:
                if line.startswith('#'):
                    header_level = len(line) - len(line.lstrip('#'))
                    header_text = line.lstrip('#').strip()
                    headers.append({
                        'level': header_level,
                        'text': header_text
                    })
            
            metadata["content"] = content
            metadata["line_count"] = len(content)
            metadata["headers"] = headers
            metadata["header_count"] = len(headers)
            
            return metadata
        
        except UnicodeDecodeError:
            # UTF-8로 디코딩 실패 시 다른 인코딩 시도
            try:
                with open(file_path, 'r', encoding='cp949') as f:
                    content = f.readlines()
                content = [line.strip() for line in content if line.strip()]
                metadata["content"] = content
                metadata["line_count"] = len(content)
                return metadata
            except Exception as e:
                metadata["content"] = ["[파싱 오류 발생]"]
                metadata["error"] = str(e)
                return metadata
        except Exception as e:
            print(f"마크다운 파일 파싱 오류 ({file_path}): {str(e)}")
            metadata["content"] = ["[파싱 오류 발생]"]
            metadata["error"] = str(e)
            return metadata
    
    def parse_file(self, file_path: str = None) -> Dict[str, Any]:
        """
        파일 유형에 따라 적절한 파서 호출
        
        Args:
            file_path (str, optional): 파일 경로
            
        Returns:
            Dict[str, Any]: 파싱 결과와 메타데이터
        """
        if file_path is None:
            file_path = self.file_path
        
        if file_path is None:
            raise ValueError("파일 경로가 지정되지 않았습니다")
        
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext == '.pdf':
            return self.parse_pdf(file_path)
        elif file_ext in ['.docx', '.doc']:
            return self.parse_docx(file_path)
        elif file_ext in ['.pptx', '.ppt']:
            return self.parse_pptx(file_path)
        elif file_ext == '.txt':
            return self.parse_txt(file_path)
        elif file_ext == '.md':
            return self.parse_markdown(file_path)
        else:
            # 지원하지 않는 파일 유형
            metadata = self._extract_metadata_from_filename(file_path)
            metadata["content"] = ["[지원하지 않는 파일 형식]"]
            metadata["error"] = f"지원하지 않는 파일 형식: {file_ext}"
            return metadata
    
    def parse_directory(self, directory: str = None) -> List[Dict[str, Any]]:
        """
        지정된 디렉토리의 모든 파일 파싱
        
        Args:
            directory (str, optional): 파싱할 디렉토리. 없으면 DATA_DIR 사용
            
        Returns:
            List[Dict[str, Any]]: 파싱 결과와 메타데이터 목록
        """
        if directory is None:
            directory = DATA_DIR
        
        result = []
        supported_exts = ['.pdf', '.docx', '.doc', '.pptx', '.ppt', '.txt', '.md']
        
        for root, _, files in os.walk(directory):
            for file in files:
                file_path = os.path.join(root, file)
                file_ext = os.path.splitext(file)[1].lower()
                
                if file_ext in supported_exts:
                    try:
                        parsed_data = self.parse_file(file_path)
                        result.append(parsed_data)
                    except Exception as e:
                        print(f"파일 '{file}' 파싱 중 오류 발생: {str(e)}")
        
        return result
    
    def split_documents(self, documents: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        문서를 청크로 분할
        
        Args:
            documents (List[Dict[str, Any]]): 파싱된 문서 목록
            
        Returns:
            List[Dict[str, Any]]: 청크로 분할된 문서 목록
        """
        chunked_documents = []
        
        for doc in documents:
            metadata = {k: v for k, v in doc.items() if k != 'content'}
            
            # 각 content 항목을 분할
            for i, content_item in enumerate(doc.get('content', [])):
                chunks = self.text_splitter.split_text(content_item)
                
                for j, chunk in enumerate(chunks):
                    chunk_metadata = metadata.copy()
                    chunk_metadata['chunk_id'] = f"{i}_{j}"
                    chunk_metadata['original_content_index'] = i
                    chunk_metadata['chunk_index'] = j
                    
                    chunked_documents.append({
                        'content': chunk,
                        'metadata': chunk_metadata
                    })
        
        return chunked_documents

def main():
    """테스트용 메인 함수"""
    parser = DocumentParser()
    parsed_documents = parser.parse_directory()
    print(f"총 {len(parsed_documents)}개 문서 파싱 완료")
    
    # 결과 저장
    with open('parsed_documents.json', 'w', encoding='utf-8') as f:
        json.dump(parsed_documents, f, ensure_ascii=False, indent=2)
    
    # 청크 분할 테스트
    chunked_docs = parser.split_documents(parsed_documents)
    print(f"총 {len(chunked_docs)}개 청크로 분할됨")
    
    # 청크 샘플 출력
    for i, doc in enumerate(chunked_docs[:5]):
        print(f"\n--- 청크 {i+1} ---")
        print(f"메타데이터: {doc['metadata']}")
        print(f"내용 일부: {doc['content'][:150]}...")

if __name__ == "__main__":
    main() 